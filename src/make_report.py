"""Genera el reporte de riesgo en Excel a partir de la base de datos y el modelo entrenado.

Hojas: Resumen, Segmentos, Deciles, Variables, Alto riesgo y una hoja por cada consulta SQL.
Se puede programar (Programador de tareas / cron) para que se genere solo cada mes.

Uso:
    python src/make_report.py
"""

import json
import sqlite3
from datetime import date

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from features import DB, MODELO, REPORTS
from run_queries import QUERIES, leer_consultas

MODEL_DIR = REPORTS / "model"
AZUL = "3C3489"
ENCABEZADO = PatternFill("solid", fgColor=AZUL)
TOP_CLIENTES = 500


def cargar_datos() -> dict[str, pd.DataFrame]:
    with sqlite3.connect(DB) as conn:
        cartera = pd.read_sql_query(
            "SELECT c.*, p.prob_impago, p.segmento_riesgo "
            "FROM clientes c JOIN predicciones p USING (cliente_id)", conn)
        consultas = {n: pd.read_sql_query(sql, conn) for n, sql in leer_consultas(QUERIES).items()}

    segmentos = cartera.groupby("segmento_riesgo").agg(
        clientes=("cliente_id", "size"),
        prob_promedio=("prob_impago", "mean"),
        tasa_impago_real=("impago_2_anios", "mean"),
        ingreso_promedio=("ingreso_mensual", "mean"),
    ).reset_index()
    segmentos.insert(2, "pct_cartera", segmentos["clientes"] / segmentos["clientes"].sum())

    alto_riesgo = (cartera.sort_values("prob_impago", ascending=False).head(TOP_CLIENTES)
                   [["cliente_id", "prob_impago", "segmento_riesgo", "edad", "ingreso_mensual",
                     "utilizacion_credito", "mora_30_59_dias", "mora_60_89_dias", "mora_90_dias"]])

    return {
        "cartera": cartera,
        "Segmentos": segmentos,
        "Deciles": pd.read_csv(MODEL_DIR / "deciles.csv"),
        "Variables": pd.read_csv(MODEL_DIR / "importancia_variables.csv"),
        "Alto riesgo": alto_riesgo,
        **{f"SQL {n}"[:31]: df for n, df in consultas.items()},
    }


def formatear_hoja(ws, df: pd.DataFrame) -> None:
    """Encabezado de color, anchos automáticos, porcentajes y filtros."""
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = ENCABEZADO
        celda.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, col in enumerate(df.columns, start=1):
        letra = get_column_letter(i)
        ancho = max(len(str(col)), *(len(f"{v:.4f}" if isinstance(v, float) else str(v)) for v in df[col].head(200)))
        ws.column_dimensions[letra].width = min(ancho + 3, 40)
        es_pct = any(p in col for p in ("prob", "tasa_impago_real", "pct_cartera", "capturados"))
        formato = "0.0%" if es_pct else ("#,##0.00" if df[col].dtype.kind == "f" else None)
        if formato:
            for (celda,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                celda.number_format = formato
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def grafico_barras(ws, titulo: str, col_cat: int, cols_val: list[int], n: int, ancla: str, pct=True):
    g = BarChart()
    g.title, g.height, g.width = titulo, 8, 16
    for c in cols_val:
        g.add_data(Reference(ws, min_col=c, min_row=1, max_row=n + 1), titles_from_data=True)
    g.set_categories(Reference(ws, min_col=col_cat, min_row=2, max_row=n + 1))
    if pct:
        g.y_axis.numFmt = "0%"
    ws.add_chart(g, ancla)


def hoja_resumen(wb, datos: dict, metricas: dict) -> None:
    ws = wb.create_sheet("Resumen", 0)
    cartera = datos["cartera"]
    filas = [
        ("Reporte de riesgo de crédito", None),
        (f"Generado el {date.today():%d/%m/%Y}", None),
        (None, None),
        ("CARTERA", None),
        ("Clientes analizados", len(cartera)),
        ("Tasa de impago observada", cartera["impago_2_anios"].mean()),
        ("Probabilidad de impago promedio (modelo)", cartera["prob_impago"].mean()),
        ("Clientes en segmento Alto o Muy alto",
         int(cartera["segmento_riesgo"].isin(["3. Alto", "4. Muy alto"]).sum())),
        (None, None),
        ("MODELO", None),
        ("Algoritmo", metricas["modelo"]),
        ("Fecha de entrenamiento", metricas["fecha_entrenamiento"]),
        ("AUC (conjunto de prueba)", metricas["auc"]),
        ("Gini", metricas["gini"]),
        ("KS", metricas["ks"]),
        ("Brier score", metricas["brier"]),
        (None, None),
        ("Cómo leerlo: AUC de 0,5 es azar y 1 es perfecto. En riesgo de crédito, "
         "un Gini superior a 0,5 se considera un buen modelo.", None),
    ]
    for fila in filas:
        ws.append(fila)
    ws["A1"].font = Font(bold=True, size=16, color=AZUL)
    ws["A2"].font = Font(italic=True, color="666666")
    for celda in ("A4", "A10"):
        ws[celda].font = Font(bold=True, color=AZUL)
    for celda in ("B6", "B7"):
        ws[celda].number_format = "0.00%"
    ws["B5"].number_format = ws["B8"].number_format = "#,##0"
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 22


def main() -> None:
    if not MODELO.exists():
        raise SystemExit("Primero ejecuta: python src/train_model.py")
    metricas = json.loads((MODEL_DIR / "metricas.json").read_text(encoding="utf-8"))
    datos = cargar_datos()
    salida = REPORTS / f"reporte_riesgo_{date.today():%Y-%m}.xlsx"

    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        for nombre, df in datos.items():
            if nombre == "cartera":
                continue
            df.to_excel(writer, sheet_name=nombre, index=False)
            formatear_hoja(writer.sheets[nombre], df)

        wb = writer.book
        hoja_resumen(wb, datos, metricas)
        n_seg, n_var = len(datos["Segmentos"]), len(datos["Variables"])
        grafico_barras(wb["Segmentos"], "Tasa de impago: predicha vs. real", 1, [4, 5], n_seg, "H2")
        grafico_barras(wb["Variables"], "Importancia de variables (caída de AUC)", 1, [2], n_var, "D2", pct=False)

        ws = wb["Deciles"]
        g = LineChart()
        g.title, g.height, g.width = "Calibración por decil de riesgo", 8, 16
        g.add_data(Reference(ws, min_col=3, max_col=4, min_row=1, max_row=11), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=1, min_row=2, max_row=11))
        g.y_axis.numFmt, g.x_axis.title = "0%", "Decil (10 = más riesgo)"
        ws.add_chart(g, "H2")

    print(f"Reporte generado: {salida}")


if __name__ == "__main__":
    main()
